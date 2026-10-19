"""Answer files key bracket decisions on anchors (spec 2026-10-13, section 2);
BRK ids stay for display."""

import json

from openpyxl import load_workbook

from specgen import secfile
from specgen.answers import read_answers_json, read_workbook, write_workbook
from specgen.htmlpage import write_html
from specgen.manifest import build_manifest

SRC = ("<SEC><PRT><TTL>PART 2 PRODUCTS</TTL><SPT><TTL>CHILLERS</TTL>"
       "<TXT>Provide [scroll][screw] compressors.</TXT></SPT></PRT></SEC>")
# a single PRT is numbered by position (PART 1), not by its title text
ANCHOR = "PART 1/CHILLERS#choice:e37dc8c8"


def _m():
    return build_manifest(secfile.parse_source(SRC))


def test_manifest_carries_anchor():
    assert _m()["decisions"][0]["anchor"] == ANCHOR


def _fill(xlsx, answer):
    wb = load_workbook(xlsx)
    ws = wb["Answers"]
    head = [c.value for c in ws[1]]
    ws.cell(row=2, column=head.index("Answer") + 1, value=answer)
    wb.save(xlsx)
    return head


def test_workbook_round_trips_anchor_keys(tmp_path):
    x = tmp_path / "a.xlsx"
    write_workbook(_m(), x)
    head = _fill(x, "1")
    assert head[-1] == "Anchor"
    _, answers, _ = read_workbook(x)
    assert answers == {ANCHOR: {"answer": "1", "notes": ""}}


def test_old_workbook_without_anchor_column_keys_on_id(tmp_path):
    x = tmp_path / "a.xlsx"
    write_workbook(_m(), x)
    wb = load_workbook(x)
    wb["Answers"].delete_cols(len(list(wb["Answers"][1])))  # drop Anchor
    wb.save(x)
    _fill(x, "1")
    _, answers, _ = read_workbook(x)
    assert answers == {"BRK-0001": {"answer": "1", "notes": ""}}


def test_json_v2_keeps_display_fields(tmp_path):
    p = tmp_path / "a.json"
    p.write_text(json.dumps({"schema": "specgen.answers/2", "answers": {
        ANCHOR: {"answer": "1", "id": "BRK-0001", "options": ["scroll", "screw"]}}}))
    _, answers, _ = read_answers_json(p)
    assert answers[ANCHOR] == {"answer": "1", "notes": "", "id": "BRK-0001",
                               "options": ["scroll", "screw"]}


def test_html_exports_anchor_keys_and_imports_them(tmp_path):
    out = tmp_path / "m.html"
    write_html(_m(), out)
    html = out.read_text(encoding="utf-8")
    assert f'"anchor": "{ANCHOR}"' in html
    assert 'schema: "specgen.answers/2"' in html
    assert "keyed[d.anchor]" in html          # export keys on the anchor
    assert "byAnchor" in html                 # import maps anchors back to ids
    assert "not in this master" in html       # orphaned imports are reported

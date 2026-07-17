"""Additive bracket lists (issues #33/#34) and inline remove-gating (#35).

UFGS adjacent bracket groups are not always exclusive-or:
``[ single pole][, double pole][, and four-way]`` is a select-any list, and
``[nonmetallic][ or ][cast-metal]`` carries a connector-only span that joins
its neighbors. The scanner classifies these, the generator recomposes kept
cores with fresh separators, and the manifest flags mid-sentence groups so
the HTML page demotes "remove entirely".
"""

import json

from specgen import secfile
from specgen.generate import generate
from specgen.htmlpage import write_html
from specgen.manifest import build_manifest
from specgen.scan import (
    additive_join,
    connector_word,
    is_additive,
    is_connector,
    scan_groups,
    strip_connector,
)

SWITCHES = (
    "<SEC><TXT>NEMA WD 1, UL 20[ single pole][, double pole][, three-way]"
    "[, and four-way] totally enclosed.</TXT></SEC>"
)
OUTLETS = (
    "<SEC><TXT>Each outlet: consisting of [nonmetallic][ or ][cast-metal]"
    " body with threaded openings.</TXT></SEC>"
)
FUSES = (
    "<SEC><TXT>Provide fuses for each fusible[ switch][ panel]"
    "[ and control center].</TXT></SEC>"
)
EXCLUSIVE = "<SEC><TXT>Provide [scroll][screw] compressors.</TXT></SEC>"


def _groups(src):
    return scan_groups(src)[0]


def test_connector_only_spans():
    assert is_connector(" or ")
    assert is_connector(", and")
    assert is_connector(",")
    assert not is_connector("orifice plate")
    assert not is_connector("cast-metal")
    assert not is_connector("   ")


def test_additive_classification():
    assert is_additive(_groups(SWITCHES)[0])
    assert is_additive(_groups(OUTLETS)[0])
    assert is_additive(_groups(FUSES)[0])
    assert not is_additive(_groups(EXCLUSIVE)[0])


def test_connector_word_comes_from_master_text():
    assert connector_word(_groups(SWITCHES)[0]) == "and"
    assert connector_word(_groups(OUTLETS)[0]) == "or"
    assert connector_word(_groups(FUSES)[0]) == "and"


def test_strip_connector_and_join():
    assert strip_connector(", double pole") == "double pole"
    assert strip_connector(" and control center") == "control center"
    assert strip_connector(" single pole") == "single pole"
    assert strip_connector("or fiberglass conduit") == "fiberglass conduit"
    # no connector present: text is only trimmed, never truncated
    assert strip_connector("Order of precedence") == "Order of precedence"
    assert additive_join(["a"], "and") == "a"
    assert additive_join(["a", "b"], "or") == "a or b"
    assert additive_join(["a", "b", "c"], "and") == "a, b, and c"


def _gen(tmp_path, src, answer, revisions=False):
    master = tmp_path / "m.SEC"
    master.write_bytes(src.encode(secfile.ENCODING))
    out = tmp_path / "out.sec"
    generate(master, {}, {"BRK-0001": {"answer": answer}}, out,
             revisions=revisions)
    return out.read_bytes().decode(secfile.ENCODING)


def test_generate_additive_subset(tmp_path):
    got = _gen(tmp_path, SWITCHES, "1,3")
    assert "UL 20 single pole and three-way totally enclosed." in got
    assert "[" not in got


def test_generate_additive_all_and_oxford_comma(tmp_path):
    got = _gen(tmp_path, FUSES, "1,2,3")
    assert "each fusible switch, panel, and control center." in got


def test_generate_additive_single_pick_strips_connector(tmp_path):
    got = _gen(tmp_path, SWITCHES, "2")
    assert "UL 20 double pole totally enclosed." in got


def test_generate_connector_span_joins_neighbors(tmp_path):
    got = _gen(tmp_path, OUTLETS, "1,3")
    assert "consisting of nonmetallic or cast-metal body" in got
    # picking the connector index explicitly changes nothing
    assert got == _gen(tmp_path, OUTLETS, "1,2,3")


def test_generate_additive_revisions_invariants(tmp_path):
    # generate enforces resolve==clean and revert==master internally and
    # raises if the DEL/ADD wrapping is wrong — success IS the assertion
    got = _gen(tmp_path, SWITCHES, "1,4", revisions=True)
    assert "single pole and four-way" in got
    rev = (tmp_path / "out.revisions.sec").read_bytes().decode(secfile.ENCODING)
    assert "<DEL>" in rev and "<ADD>" in rev


LEAD_CONN = (
    "<SEC><TXT>Provide lubrication for pumps[ and][ motors].</TXT></SEC>"
)
TRAIL_COMMAS = (
    "<SEC><TXT>Cover plate of [brass,][ stainless steel,][ or]"
    "[ high-impact plastic] construction.</TXT></SEC>"
)


def test_edge_connector_span_kept_with_cores(tmp_path):
    # "[ and][ motors]": the leading connector joins the group to text
    # OUTSIDE it ("pumps") — it must be kept when a core is kept
    got = _gen(tmp_path, LEAD_CONN, "2")
    assert "for pumps and motors." in got
    # nothing picked but the connector alone = nothing kept
    from specgen.scan import connector_runs
    g = _groups(LEAD_CONN)[0]
    assert connector_runs(g) == (1, 2)


def test_trailing_commas_do_not_double(tmp_path):
    got = _gen(tmp_path, TRAIL_COMMAS, "1,2,4")
    assert "of brass, stainless steel, or high-impact plastic construction." in got
    got2 = _gen(tmp_path, TRAIL_COMMAS, "2")
    assert "of stainless steel construction." in got2


def test_exclusive_choice_unchanged(tmp_path):
    got = _gen(tmp_path, EXCLUSIVE, "2")
    assert "Provide screw compressors." in got


def _manifest(tmp_path, src):
    p = tmp_path / "m.SEC"
    p.write_bytes(src.encode(secfile.ENCODING))
    return build_manifest(secfile.parse(p))


def test_manifest_additive_fields(tmp_path):
    d = _manifest(tmp_path, OUTLETS)["decisions"][0]
    assert d["additive"] is True
    assert d["connector"] == "or"
    assert d["conn_opts"] == [1]
    e = _manifest(tmp_path, EXCLUSIVE)["decisions"][0]
    assert e["additive"] is False and e["conn_opts"] == []


def test_manifest_inline_flag(tmp_path):
    # mid-sentence: inline
    assert _manifest(tmp_path, FUSES)["decisions"][0]["inline"] is True
    # a group alone on its line(s): not inline
    block = "<SEC><TXT>Intro.\n[Optional paragraph text.]\nNext.</TXT></SEC>"
    assert _manifest(tmp_path, block)["decisions"][0]["inline"] is False


def test_html_page_carries_additive_ui(tmp_path):
    p = tmp_path / "m.SEC"
    p.write_bytes(SWITCHES.encode(secfile.ENCODING))
    out = tmp_path / "m.html"
    write_html(build_manifest(secfile.parse(p)), out)
    html = out.read_text(encoding="utf-8")
    assert "additiveJoin" in html and "additiveCore" in html
    assert "removelink" in html          # demoted remove-entirely style
    assert "rarely correct" in html      # demotion copy
    # the embedded manifest carries the new flags
    assert '"additive": true' in html and '"connector": "and"' in html


def test_workbook_marks_additive_select_all(tmp_path):
    from openpyxl import load_workbook

    from specgen.answers import write_workbook

    m = _manifest(tmp_path, SWITCHES)
    xlsx = tmp_path / "a.xlsx"
    write_workbook(m, xlsx)
    ws = load_workbook(xlsx)["Answers"]
    opts_cell = ws.cell(row=2, column=6).value
    assert opts_cell.startswith("SELECT ALL THAT APPLY")

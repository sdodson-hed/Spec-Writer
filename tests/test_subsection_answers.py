"""Subsection answers: JSON, workbook, and resolution (spec 2026-10-13,
section 3; plan 2026-10-19 Task 2, decision D1 for governed subsections)."""

import json

import pytest
from openpyxl import load_workbook

from specgen import secfile
from specgen.answers import (
    read_answers_json,
    read_workbook,
    resolve_answer_keys,
    write_workbook,
)
from specgen.generate import GenerateError, generate
from specgen.manifest import build_manifest
from tests.test_subsections import SRC


def _m(src=SRC):
    return build_manifest(secfile.parse_source(src))


def _subs(m):
    return {s["id"]: s for s in m["subsections"]}


def _resolve(answers, src=SRC, fingerprint=None, **kw):
    from specgen.anchors import section_anchors
    from specgen.scan import scan_groups

    sec = secfile.parse_source(src)
    return resolve_answer_keys(
        answers, scan_groups(src)[0], section_anchors(sec), fingerprint,
        subsections=build_manifest(sec)["subsections"], **kw,
    )


def _a(m, sid, answer, **extra):
    s = _subs(m)[sid]
    return {s["anchor"]: {"answer": answer, "id": sid, "num": s["num"],
                          "title": s["title"], **extra}}


# --- JSON and workbook -----------------------------------------------------

def test_json_keeps_subsection_display_fields(tmp_path):
    m = _m()
    p = tmp_path / "a.json"
    p.write_text(json.dumps({"schema": "specgen.answers/2",
                             "answers": _a(m, "SPT-0003", "remove")}))
    _, answers, _ = read_answers_json(p)
    [(key, rec)] = answers.items()
    assert key.endswith("#spt:da4313f4")
    assert rec == {"answer": "remove", "notes": "", "id": "SPT-0003",
                   "num": "2.1.1", "title": "Scroll"}


def test_workbook_subsections_sheet_round_trip(tmp_path):
    m = _m()
    x = tmp_path / "a.xlsx"
    write_workbook(m, x)
    wb = load_workbook(x)
    ws = wb["Subsections"]
    head = [c.value for c in ws[1]]
    assert head == ["ID", "Number", "Title", "Governed by", "Action", "Anchor"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["SPT-0004"][3].value == _subs(m)["SPT-0004"]["governed_by"]
    rows["SPT-0003"][4].value = "remove"
    rows["SPT-0007"][4].value = "keep"
    wb.save(x)
    _, answers, _ = read_workbook(x)
    s = _subs(m)
    assert answers[s["SPT-0003"]["anchor"]] == {"answer": "remove", "notes": ""}
    assert answers[s["SPT-0007"]["anchor"]] == {"answer": "keep", "notes": ""}
    assert len(answers) == 2


def test_old_workbook_without_subsections_sheet_loads(tmp_path):
    x = tmp_path / "a.xlsx"
    write_workbook(_m(), x)
    wb = load_workbook(x)
    del wb["Subsections"]
    wb.save(x)
    assert read_workbook(x)[1] == {}


# --- resolution --------------------------------------------------------------

def test_subsection_anchor_resolves_to_its_id():
    m = _m()
    by_gid, orphans, status = _resolve(_a(m, "SPT-0003", "remove"))
    assert by_gid["SPT-0003"]["answer"] == "remove"
    assert orphans == [] and status == "anchors"


def test_subsection_answer_values_are_normalized_and_checked():
    m = _m()
    by_gid, problems, _ = _resolve({**_a(m, "SPT-0003", " Remove "),
                                    **_a(m, "SPT-0007", "maybe")})
    assert by_gid["SPT-0003"]["answer"] == "remove"
    assert "SPT-0007" not in by_gid
    [p] = problems
    assert p["id"] == "SPT-0007" and "keep or remove" in p["note"]


def test_subsection_orphan_reports_old_title():
    m = _m()
    renamed = SRC.replace("<TTL>Scroll</TTL>", "<TTL>Scroll Compressors</TTL>")
    by_gid, orphans, _ = _resolve(_a(m, "SPT-0003", "remove"), src=renamed)
    assert "SPT-0003" not in by_gid
    [o] = orphans
    assert o["options"] == ["2.1.1 Scroll"]
    assert "master changed: decision not found (was SPT-0003)" in o["note"]


def test_positional_subsection_keys_are_rejected():
    with pytest.raises(GenerateError, match="keyed by anchor"):
        _resolve({"SPT-0003": {"answer": "remove"}})


def test_strict_counts_subsection_problems():
    m = _m()
    with pytest.raises(GenerateError, match="1 answer"):
        _resolve(_a(m, "SPT-0007", "maybe"), strict=True)


# --- governed subsections (D1) ---------------------------------------------

def test_governed_remove_becomes_bracket_delete():
    m = _m()
    brk = _subs(m)["SPT-0004"]["governed_by"]
    by_gid, problems, _ = _resolve(_a(m, "SPT-0004", "remove", notes="no centrifugal"))
    assert "SPT-0004" not in by_gid
    assert by_gid[brk] == {"answer": "delete", "notes": "no centrifugal", "via": "SPT-0004"}
    assert problems == []


def test_governed_keep_becomes_bracket_keep():
    m = _m()
    brk = _subs(m)["SPT-0004"]["governed_by"]
    by_gid, _, _ = _resolve(_a(m, "SPT-0004", "keep"))
    assert by_gid[brk]["answer"] == "keep"


def test_governed_agreement_is_fine():
    m = _m()
    s = _subs(m)["SPT-0004"]
    brk_anchor = next(d["anchor"] for d in m["decisions"] if d["id"] == s["governed_by"])
    by_gid, problems, _ = _resolve({**_a(m, "SPT-0004", "remove"),
                                    brk_anchor: {"answer": "Delete"}})
    assert by_gid[s["governed_by"]]["answer"] == "Delete" and problems == []


def test_governed_conflict_applies_neither():
    m = _m()
    s = _subs(m)["SPT-0004"]
    brk_anchor = next(d["anchor"] for d in m["decisions"] if d["id"] == s["governed_by"])
    by_gid, _, _ = _resolve({**_a(m, "SPT-0004", "remove"),
                             brk_anchor: {"answer": "keep"}})
    rec = by_gid[s["governed_by"]]
    assert rec["answer"] == ""
    assert "conflicting answers" in rec["conflict"]
    assert "SPT-0004" in rec["conflict"] and "keep" in rec["conflict"]


# --- generate --------------------------------------------------------------

def _gen(tmp_path, answers):
    p = tmp_path / "m.SEC"
    p.write_bytes(SRC.encode(secfile.ENCODING))
    s = generate(p, {}, answers, tmp_path / "o.sec")
    out = (tmp_path / "o.sec").read_bytes().decode(secfile.ENCODING)
    redline = json.loads((tmp_path / "o.redline.json").read_text(encoding="utf-8"))
    unresolved = (tmp_path / "o.unresolved.md").read_text(encoding="utf-8")
    return s, out, redline, unresolved


def test_generate_applies_governed_removal(tmp_path):
    m = _m()
    _, out, redline, _ = _gen(tmp_path, _a(m, "SPT-0004", "remove"))
    assert "Centrifugal" not in out and "Screw" in out
    brk = _subs(m)["SPT-0004"]["governed_by"]
    [c] = [c for c in redline["changes"] if c["id"] == brk]
    assert c["action"] == "deleted (via SPT-0004)"


def test_generate_conflict_leaves_bracket_unresolved(tmp_path):
    m = _m()
    s = _subs(m)["SPT-0004"]
    brk_anchor = next(d["anchor"] for d in m["decisions"] if d["id"] == s["governed_by"])
    _, out, _, unresolved = _gen(tmp_path, {**_a(m, "SPT-0004", "remove"),
                                            brk_anchor: {"answer": "keep"}})
    assert "[<SPT><TTL>Centrifugal" in out
    assert "conflicting answers" in unresolved


def test_generate_accepts_subsection_answers(tmp_path):
    m = _m()
    summary, _, _, _ = _gen(tmp_path, _a(m, "SPT-0003", "remove"))
    assert summary["answers_orphaned"] == 0

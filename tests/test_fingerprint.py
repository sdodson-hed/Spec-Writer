"""Answers fingerprint (issue #23): exports are stamped with a hash of the
decision surface, and generate refuses answer files recorded against a
different scan — a master edit or scanner change can silently shift what a
BRK id points at, which is how a stale file writes the wrong spec.
"""

import json

import pytest
from openpyxl import load_workbook

from specgen import secfile
from specgen.answers import read_answers_json, read_workbook, write_workbook
from specgen.generate import GenerateError, generate
from specgen.htmlpage import write_html
from specgen.manifest import build_manifest
from specgen.scan import legacy_scan_fingerprint, scan_fingerprint, scan_groups

SRC = "<SEC><TXT>Provide [scroll][screw] compressors.</TXT></SEC>"
SRC2 = "<SEC><TXT>Extra [text]. Provide [scroll][screw] compressors.</TXT></SEC>"


def _manifest(tmp_path, src=SRC, name="m.SEC"):
    p = tmp_path / name
    p.write_bytes(src.encode(secfile.ENCODING))
    return build_manifest(secfile.parse(p))


def test_fingerprint_tracks_decision_surface(tmp_path):
    fp1 = scan_fingerprint(scan_groups(SRC)[0])
    fp2 = scan_fingerprint(scan_groups(SRC2)[0])
    assert fp1 != fp2  # ids shifted: BRK-0001 means different brackets
    assert fp1 == scan_fingerprint(scan_groups(SRC)[0])  # deterministic
    assert _manifest(tmp_path)["fingerprint"] == fp1


def test_workbook_meta_roundtrip(tmp_path):
    m = _manifest(tmp_path)
    xlsx = tmp_path / "a.xlsx"
    write_workbook(m, xlsx)
    assert "Meta" in load_workbook(xlsx).sheetnames
    _, _, fp = read_workbook(xlsx)
    assert fp == m["fingerprint"]


def test_answers_json_carries_fingerprint(tmp_path):
    p = tmp_path / "a.json"
    p.write_text(json.dumps({"fingerprint": "abc123", "answers": {}}))
    _, _, fp = read_answers_json(p)
    assert fp == "abc123"
    p.write_text(json.dumps({"answers": {}}))  # pre-fingerprint export
    assert read_answers_json(p)[2] is None


def test_html_export_stamps_fingerprint(tmp_path):
    m = _manifest(tmp_path)
    out = tmp_path / "m.html"
    write_html(m, out)
    html = out.read_text(encoding="utf-8")
    assert "fingerprint: DATA.fingerprint" in html
    assert f'"fingerprint": "{m["fingerprint"]}"' in html


def _gen(tmp_path, fingerprint, force_stale=False):
    master = tmp_path / "g.SEC"
    master.write_bytes(SRC.encode(secfile.ENCODING))
    return generate(
        master,
        {},
        {"BRK-0001": {"answer": "1"}},
        tmp_path / "out.sec",
        fingerprint=fingerprint,
        force_stale=force_stale,
    )


def test_generate_refuses_stale_fingerprint(tmp_path):
    good = scan_fingerprint(scan_groups(SRC)[0])
    stale = scan_fingerprint(scan_groups(SRC2)[0])
    with pytest.raises(GenerateError, match="different scan"):
        _gen(tmp_path, stale)
    assert _gen(tmp_path, good)["answers_fingerprint"] == "match"
    assert _gen(tmp_path, None)["answers_fingerprint"] == "unstamped"
    forced = _gen(tmp_path, stale, force_stale=True)
    assert forced["answers_fingerprint"] == "MISMATCH (forced)"


CRLF = "<SEC><TXT>Provide [scroll\r\n][screw] compressors.</TXT>\r\n</SEC>"
LF = "<SEC><TXT>Provide [scroll\n][screw] compressors.</TXT>\n</SEC>"
REFLOW = "<SEC><TXT>Longer text before. Provide [scroll\n][screw] compressors.</TXT>\n</SEC>"


def test_fingerprint_ignores_line_endings():
    assert scan_fingerprint(scan_groups(CRLF)[0]) == scan_fingerprint(scan_groups(LF)[0])


def test_fingerprint_ignores_text_outside_brackets():
    # 23 64 10 Apr->Jul 2026: reference article reflowed, no bracket changed
    assert scan_fingerprint(scan_groups(LF)[0]) == scan_fingerprint(scan_groups(REFLOW)[0])


def test_fingerprint_still_tracks_option_text_and_ids():
    other = LF.replace("screw", "rotary screw")
    assert scan_fingerprint(scan_groups(LF)[0]) != scan_fingerprint(scan_groups(other)[0])
    # SRC2 inserts a group before: BRK-0001 now means something else
    assert scan_fingerprint(scan_groups(SRC)[0]) != scan_fingerprint(scan_groups(SRC2)[0])


def test_generate_accepts_legacy_fingerprint(tmp_path):
    legacy = legacy_scan_fingerprint(scan_groups(SRC)[0])
    assert legacy != scan_fingerprint(scan_groups(SRC)[0])
    assert _gen(tmp_path, legacy)["answers_fingerprint"] == "match (legacy)"
